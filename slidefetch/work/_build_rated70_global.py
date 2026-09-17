#!/usr/bin/env python3
"""Build a global DB for the rated-70 slide download + a per-folder feedback .xlsx.

Global DB  : rated70/rated70.db
  - courses   (folder PK 01..N) : course + slide link + RMP prof + ratings summary + folder/excel paths
  - professors                  : RMP professor match(es) per course (from ccr_rmp.db)
  - ratings                     : individual RMP ratings per course (from ccr_rmp.db)
  - feedback  (folder PK)       : review fields, seeded from ccr_rmp.db feedback, + excel_path

Per-folder   : rated70/NN/NN_feedback.xlsx   (Course / Feedback / Decks / RMP_ratings sheets)
Link         : global DB stores each folder's excel_path + folder_path; each xlsx carries the
               folder + course_college key back to the DB and a hyperlink to the slide_link.
"""
from __future__ import annotations

import csv
import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path("/home/b-ysonale/slidefetch/rated70")
GLOBAL_DB = ROOT / "rated70.db"
CCR = "/home/b-ysonale/ccr_rmp/ccr_rmp.db"

COURSE_COLS = [
    "instructor", "rmp_status", "rmp_matched_name", "rmp_profile_url",
    "rmp_avg_rating_overall", "rmp_avg_difficulty_overall", "rmp_num_ratings_overall",
    "course_num_ratings", "course_avg_quality", "course_avg_difficulty",
    "course_code", "college_name",
]
FB_FIELDS = ["instructor_ok", "instructor_correction", "rmp_ok", "slides_ok",
             "notes", "needs_review"]


def load_manifest():
    rows = []
    with open(ROOT / "manifest.csv", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            rows.append(r)
    return rows


def deck_files(folder):
    d = ROOT / folder
    if not d.is_dir():
        return []
    return sorted(p.name for p in d.iterdir() if p.is_file() and not p.name.endswith(".xlsx"))


def init_db():
    if GLOBAL_DB.exists():
        GLOBAL_DB.unlink()
    con = sqlite3.connect(GLOBAL_DB)
    con.execute("""CREATE TABLE courses (
        folder TEXT PRIMARY KEY, course_college TEXT, course_code TEXT, college_name TEXT,
        course_date INTEGER, within3 INTEGER, slide_link TEXT, num_decks INTEGER,
        deck_files TEXT, instructor TEXT, rmp_status TEXT, rmp_matched_name TEXT,
        rmp_profile_url TEXT, rmp_avg_rating REAL, rmp_avg_difficulty REAL,
        rmp_num_ratings INTEGER, course_num_ratings INTEGER, course_avg_quality REAL,
        course_avg_difficulty REAL, folder_path TEXT, excel_path TEXT)""")
    con.execute("""CREATE TABLE professors (
        folder TEXT, course_college TEXT, instructor_name TEXT, rmp_status TEXT,
        matched_name TEXT, legacy_id TEXT, profile_url TEXT, department TEXT,
        avg_rating_overall REAL, avg_difficulty_overall REAL, num_ratings_overall INTEGER,
        course_num_ratings INTEGER, course_avg_quality REAL, course_avg_difficulty REAL)""")
    con.execute("""CREATE TABLE ratings (
        folder TEXT, course_college TEXT, legacy_id TEXT, class TEXT, quality REAL,
        difficulty REAL, would_take_again TEXT, grade TEXT, date TEXT, comment TEXT, tags TEXT)""")
    con.execute("""CREATE TABLE feedback (
        folder TEXT PRIMARY KEY, course_college TEXT, instructor_ok TEXT,
        instructor_correction TEXT, rmp_ok TEXT, slides_ok TEXT, notes TEXT,
        needs_review INTEGER, excel_path TEXT, updated_at TEXT)""")
    con.commit()
    return con


def write_xlsx(path, meta, decks, profs, ratings, fb):
    wb = Workbook()
    bold = Font(bold=True)
    # Course sheet (key/value)
    ws = wb.active
    ws.title = "Course"
    ws.append(["field", "value"])
    for c in ws[1]:
        c.font = bold
    for k, v in meta:
        ws.append([k, v])
        if k == "slide_link" and v:
            cell = ws.cell(row=ws.max_row, column=2)
            cell.hyperlink = v
            cell.font = Font(color="0563C1", underline="single")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90

    # Feedback sheet (fill-in)
    wf = wb.create_sheet("Feedback")
    wf.append(["folder", "course_college"] + FB_FIELDS)
    for c in wf[1]:
        c.font = bold
    wf.append([meta_dict(meta)["folder"], meta_dict(meta)["course_college"]] +
              [fb.get(k, "") for k in FB_FIELDS])
    for i, w in enumerate([8, 40] + [14, 24, 10, 10, 50, 12], 1):
        wf.column_dimensions[chr(64 + i)].width = w

    # Decks sheet
    wd = wb.create_sheet("Decks")
    wd.append(["#", "filename"])
    for c in wd[1]:
        c.font = bold
    for i, name in enumerate(decks, 1):
        wd.append([i, name])
    wd.column_dimensions["B"].width = 60

    # RMP ratings sheet
    wr = wb.create_sheet("RMP_ratings")
    wr.append(["class", "quality", "difficulty", "date", "grade", "comment"])
    for c in wr[1]:
        c.font = bold
    for r in ratings:
        wr.append([r["class"], r["quality"], r["difficulty"], r["date"], r["grade"],
                   (r["comment"] or "")[:500]])
    wr.column_dimensions["F"].width = 90
    wb.save(path)


def meta_dict(meta):
    return {k: v for k, v in meta}


def main():
    con = init_db()
    ccr = sqlite3.connect(f"file:{CCR}?mode=ro", uri=True)
    manifest = load_manifest()
    n_courses = n_decks = 0
    for row in manifest:
        folder = row["folder"]
        cc = row["course_college"]
        fdir = ROOT / folder
        decks = deck_files(folder)
        n_decks += len(decks)
        # course info from ccr_rmp.db
        cur = ccr.execute(f"SELECT {','.join(COURSE_COLS)} FROM courses WHERE course_college=?", (cc,))
        crow = cur.fetchone()
        info = dict(zip(COURSE_COLS, crow)) if crow else {k: None for k in COURSE_COLS}
        excel_rel = f"{folder}/{folder}_feedback.xlsx"
        excel_abs = str(fdir / f"{folder}_feedback.xlsx")
        con.execute("INSERT INTO courses VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (
            folder, cc, info["course_code"], info["college_name"],
            int(row["course_date"]) if row["course_date"] else None,
            int(row["within3"]) if row["within3"] else None, row["slide_link"], len(decks),
            json.dumps(decks), info["instructor"], info["rmp_status"], info["rmp_matched_name"],
            info["rmp_profile_url"], info["rmp_avg_rating_overall"], info["rmp_avg_difficulty_overall"],
            info["rmp_num_ratings_overall"], info["course_num_ratings"], info["course_avg_quality"],
            info["course_avg_difficulty"], str(fdir), excel_rel))
        # professors
        profs = [dict(zip([d[0] for d in ccr.execute("SELECT * FROM course_professors LIMIT 0").description], r))
                 for r in ccr.execute("SELECT * FROM course_professors WHERE course_college=?", (cc,))]
        for p in profs:
            con.execute("INSERT INTO professors VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (
                folder, cc, p.get("instructor_name"), p.get("rmp_status"), p.get("matched_name"),
                p.get("legacy_id"), p.get("profile_url"), p.get("department"),
                p.get("avg_rating_overall"), p.get("avg_difficulty_overall"),
                p.get("num_ratings_overall"), p.get("course_num_ratings"),
                p.get("course_avg_quality"), p.get("course_avg_difficulty")))
        # ratings
        rcols = [d[0] for d in ccr.execute("SELECT * FROM course_ratings_raw LIMIT 0").description]
        ratings = [dict(zip(rcols, r)) for r in
                   ccr.execute("SELECT * FROM course_ratings_raw WHERE course_college=?", (cc,))]
        for r in ratings:
            con.execute("INSERT INTO ratings VALUES (?,?,?,?,?,?,?,?,?,?,?)", (
                folder, cc, r.get("legacy_id"), r.get("class"), r.get("quality"), r.get("difficulty"),
                r.get("would_take_again"), r.get("grade"), r.get("date"), r.get("comment"), r.get("tags")))
        # feedback seed from ccr_rmp.db
        fb = {}
        fr = ccr.execute("SELECT instructor_ok,instructor_correction,rmp_ok,notes,needs_review "
                         "FROM feedback WHERE course_college=?", (cc,)).fetchone()
        if fr:
            fb = {"instructor_ok": fr[0], "instructor_correction": fr[1], "rmp_ok": fr[2],
                  "notes": fr[3], "needs_review": fr[4]}
        con.execute("INSERT INTO feedback VALUES (?,?,?,?,?,?,?,?,?,?)", (
            folder, cc, fb.get("instructor_ok"), fb.get("instructor_correction"), fb.get("rmp_ok"),
            fb.get("slides_ok"), fb.get("notes"), fb.get("needs_review"), excel_rel, None))

        # per-folder excel
        fdir.mkdir(parents=True, exist_ok=True)
        meta = [
            ("folder", folder), ("course_college", cc), ("course_code", info["course_code"]),
            ("college_name", info["college_name"]), ("course_date", row["course_date"]),
            ("within3_ratings", row["within3"]), ("slide_link", row["slide_link"]),
            ("num_decks", len(decks)), ("instructor", info["instructor"]),
            ("rmp_matched_name", info["rmp_matched_name"]), ("rmp_profile_url", info["rmp_profile_url"]),
            ("rmp_avg_rating", info["rmp_avg_rating_overall"]),
            ("course_num_ratings", info["course_num_ratings"]),
            ("course_avg_quality", info["course_avg_quality"]),
            ("course_avg_difficulty", info["course_avg_difficulty"]),
            ("folder_path", str(fdir)),
        ]
        write_xlsx(fdir / f"{folder}_feedback.xlsx", meta, decks, profs, ratings, fb)
        n_courses += 1

    con.commit()
    con.close()
    ccr.close()
    print(f"global DB: {GLOBAL_DB}  ({n_courses} courses, {n_decks} decks)")
    print(f"per-folder feedback .xlsx written into rated70/NN/NN_feedback.xlsx")


if __name__ == "__main__":
    main()
